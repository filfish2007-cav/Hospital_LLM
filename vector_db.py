import dotenv
import os
import json
import docx                                    # pip install python-docx
import pymupdf
import re
from uuid import uuid4

from langchain_google_genai import GoogleGenerativeAIEmbeddings
from langchain_pinecone import PineconeVectorStore
from langchain_core.documents import Document
from uuid import uuid4
from pinecone import ServerlessSpec
from pinecone import Pinecone

# завантажити дані з .env
dotenv.load_dotenv()

api_key = os.getenv("GEMINI_API_KEY")
pinecone_api_key = os.getenv("PINECONE_API_KEY")

# модель для преведення текстів у вектори(набір чисел)
embedding = GoogleGenerativeAIEmbeddings(
    model="gemini-embedding-001",
    api_key=api_key,
)


# векторна база даних
pc = Pinecone(api_key=pinecone_api_key)

# створення бази даних

index_name = "hospital-docs"  # назва бази даних

if not pc.has_index(index_name):
    pc.create_index(
        name=index_name,
        dimension=3072,    # кількість чисел у векторі
        metric="cosine",   # формула для пошуку схожих текстів
        spec=ServerlessSpec(
            cloud="aws",        # хмарна платформа(амазон)
            region="us-east-1"  # регіон
        ),
    )

index = pc.Index(index_name)

vector_store = PineconeVectorStore(
    index=index,          # база даних
    embedding=embedding   # модель для кодування
)

files = [
    ("data/hospital_docs/general.pdf", 5),
    ("data/hospital_docs/for_workers.docx", 12),   # <- put your docx name here
]

json_path = "data/hospital_docs/huge_file_ids.json"


# ---------------------------------------------------------
# 1. read the file as plain text (pdf or docx)
# ---------------------------------------------------------
def read_text(file_path):
    if file_path.endswith(".pdf"):
        pdf = pymupdf.open(file_path)
        text = "\n".join(page.get_text() for page in pdf)
    else:
        word = docx.Document(file_path)
        text = "\n".join(p.text for p in word.paragraphs)

    # invisible zero-width spaces break heading detection
    return text.replace("\u200b", "")


# ---------------------------------------------------------
# 2. split into blocks
#
# (?m)	^ будет означать «начало любой строки», а не только начало всего текста
# ^	начало строки
# (?= ... )	«проверь, что дальше идёт вот это, но не забирай это». Поэтому разрез делается перед заголовком, и заголовок остаётся в своём блоке
# #*	ноль или больше символов # (на случай ### 2. ...)
# [ \t]*	необязательные пробелы или табы
# \d+	одна или больше цифр
# \.	точка
# [ \t]+	хотя бы один пробел
# \S	любой непробельный символ, то есть начало названия
# ---------------------------------------------------------
def split_into_blocks(text, n_sections):
    blocks = re.split(r"(?m)^(?=#*[ \t]*\d+\.[ \t]+\S)", text)

    blocks = [
        block.strip()
        for block in blocks
        if block.strip()
    ]

    return blocks[-n_sections:]


# ---------------------------------------------------------
# 3. build documents (same as before)
# ---------------------------------------------------------
docs = []

for file_path, n_sections in files:
    file_name = os.path.basename(file_path)

    text = read_text(file_path)
    blocks = split_into_blocks(text, n_sections)

    print(file_name, "- кількість блоків:", len(blocks))

    for block in blocks:
        lines = block.splitlines()

        block_name = lines[0].strip().lstrip("# ")

        print("   ", block_name)   # quick check that the cut is correct

        doc = Document(
            page_content=block,
            metadata={
                "file_name": file_name,
                "block_name": block_name
            }
        )

        docs.append(doc)

# ---------------------------------------------------------
# 4. ids + upload to Pinecone (same as before)
# ---------------------------------------------------------
ids = [
    str(uuid4())
    for _ in range(len(docs))
]

print("Кількість створених ID:", len(ids))

vector_store.add_documents(
    documents=docs,
    ids=ids
)

print("Документи додані в Pinecone!")

# ---------------------------------------------------------
# 5. save ids to json (same as before)
# ---------------------------------------------------------
json_data = []

for doc, doc_id in zip(docs, ids):
    json_data.append({
        "id": doc_id,
        "file_name": doc.metadata["file_name"],
        "block_name": doc.metadata["block_name"]
    })

with open(json_path, "w", encoding="utf-8") as file:
    json.dump(
        json_data,
        file,
        ensure_ascii=False,
        indent=4
    )

print("ID збережені у:", json_path)